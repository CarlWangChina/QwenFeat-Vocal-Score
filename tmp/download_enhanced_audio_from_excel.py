#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

import requests
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

try:
    from tqdm import tqdm
except Exception:  # pragma: no cover - fallback for minimal environments
    tqdm = None


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_EXCEL = (
    REPO_ROOT
    / "tmp"
    / "把高分段的原唱-在KTV房间录音-并且加入人为说话-模拟KTV数据中有人聊天的问题.xlsx"
)
DEFAULT_OUT_DIR = Path("/cache/cqy_cache/audio-score/2000_enhanced")
DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0 Safari/537.36"
)


@dataclass(frozen=True)
class DownloadTask:
    row: int
    audio_id: str
    url: str
    output_path: Path


@dataclass(frozen=True)
class DownloadResult:
    row: int
    audio_id: str
    url: str
    output_path: str
    status: str
    bytes_written: int
    error: str = ""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download audio files from the enhanced-audio Excel sheet."
    )
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    parser.add_argument("--sheet", default=None, help="Sheet name. Defaults to active sheet.")
    parser.add_argument("--audio-id-col", default="B", help="Column containing audio_id.")
    parser.add_argument("--url-col", default="J", help="Column containing the audio URL.")
    parser.add_argument("--workers", type=int, default=8)
    parser.add_argument("--timeout", type=float, default=30.0)
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Download again even when the destination file already exists.",
    )
    parser.add_argument(
        "--no-extension",
        action="store_true",
        help="Save files as the bare audio_id without keeping the URL extension.",
    )
    parser.add_argument(
        "--duplicate-policy",
        choices=["suffix-row", "skip"],
        default="suffix-row",
        help=(
            "How to handle repeated audio_id filenames. "
            "suffix-row keeps every URL by saving duplicates as audio_id__rowN.ext; "
            "skip keeps only the first file for each audio_id."
        ),
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        default=None,
        help="CSV path for download results. Defaults to OUT_DIR/download_manifest.csv.",
    )
    parser.add_argument(
        "--failed-manifest",
        type=Path,
        default=None,
        help="CSV path for failed downloads. Defaults to OUT_DIR/failed_downloads.csv.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Only read the Excel file and print planned downloads.",
    )
    return parser.parse_args()


def cell_value(row_values: tuple[Any, ...], one_based_col: int) -> Any:
    index = one_based_col - 1
    if index >= len(row_values):
        return None
    return row_values[index]


def clean_audio_id(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return re.sub(r"[^0-9A-Za-z._-]+", "_", text)


def audio_id_from_url(url: str) -> str:
    last_part = Path(unquote(urlparse(url).path)).name
    match = re.match(r"^(\d{6})_", last_part)
    return match.group(1) if match else ""


def suffix_from_url(url: str) -> str:
    suffix = Path(unquote(urlparse(url).path)).suffix
    if not suffix or len(suffix) > 12:
        return ".m4a"
    return suffix


def load_tasks(args: argparse.Namespace) -> list[DownloadTask]:
    if not args.excel.exists():
        raise FileNotFoundError(f"Excel file not found: {args.excel}")

    wb = load_workbook(args.excel, read_only=True, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active
    audio_id_col = column_index_from_string(args.audio_id_col)
    url_col = column_index_from_string(args.url_col)

    tasks: list[DownloadTask] = []
    seen_outputs: set[Path] = set()
    skipped_empty = 0
    skipped_duplicate = 0
    renamed_duplicate = 0

    for row_number, row_values in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        raw_url = cell_value(row_values, url_col)
        url = str(raw_url).strip() if raw_url is not None else ""
        if not url:
            skipped_empty += 1
            continue

        audio_id = clean_audio_id(cell_value(row_values, audio_id_col))
        if not audio_id:
            audio_id = audio_id_from_url(url)
        if not audio_id:
            skipped_empty += 1
            continue

        suffix = "" if args.no_extension else suffix_from_url(url)
        output_path = args.out_dir / f"{audio_id}{suffix}"
        if output_path in seen_outputs:
            if args.duplicate_policy == "skip":
                skipped_duplicate += 1
                continue
            output_path = args.out_dir / f"{audio_id}__row{row_number}{suffix}"
            renamed_duplicate += 1
        seen_outputs.add(output_path)
        tasks.append(
            DownloadTask(
                row=row_number,
                audio_id=audio_id,
                url=url,
                output_path=output_path,
            )
        )

    if skipped_empty:
        print(f"Skipped {skipped_empty} rows with empty URL/audio_id.")
    if skipped_duplicate:
        print(f"Skipped {skipped_duplicate} duplicate output filenames.")
    if renamed_duplicate:
        print(
            "Renamed "
            f"{renamed_duplicate} duplicate output filenames with __row suffixes."
        )
    return tasks


def download_one(
    task: DownloadTask,
    *,
    timeout: float,
    retries: int,
    overwrite: bool,
) -> DownloadResult:
    if task.output_path.exists() and task.output_path.stat().st_size > 0 and not overwrite:
        return DownloadResult(
            row=task.row,
            audio_id=task.audio_id,
            url=task.url,
            output_path=str(task.output_path),
            status="exists",
            bytes_written=task.output_path.stat().st_size,
        )

    task.output_path.parent.mkdir(parents=True, exist_ok=True)
    part_path = task.output_path.with_name(
        f"{task.output_path.name}.part-row{task.row}-pid{os.getpid()}"
    )
    headers = {"User-Agent": DEFAULT_USER_AGENT}
    last_error = ""

    for attempt in range(1, retries + 1):
        bytes_written = 0
        try:
            with requests.get(
                task.url,
                headers=headers,
                stream=True,
                timeout=(10.0, timeout),
            ) as response:
                response.raise_for_status()
                with part_path.open("wb") as handle:
                    for chunk in response.iter_content(chunk_size=1024 * 256):
                        if not chunk:
                            continue
                        handle.write(chunk)
                        bytes_written += len(chunk)

            if bytes_written <= 0:
                raise RuntimeError("downloaded zero bytes")

            part_path.replace(task.output_path)
            return DownloadResult(
                row=task.row,
                audio_id=task.audio_id,
                url=task.url,
                output_path=str(task.output_path),
                status="downloaded",
                bytes_written=bytes_written,
            )
        except Exception as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            try:
                part_path.unlink()
            except FileNotFoundError:
                pass
            if attempt < retries:
                time.sleep(min(2 ** (attempt - 1), 8))

    return DownloadResult(
        row=task.row,
        audio_id=task.audio_id,
        url=task.url,
        output_path=str(task.output_path),
        status="failed",
        bytes_written=0,
        error=last_error,
    )


def write_manifest(path: Path, results: list[DownloadResult]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "row",
                "audio_id",
                "url",
                "output_path",
                "status",
                "bytes_written",
                "error",
            ],
        )
        writer.writeheader()
        for result in sorted(results, key=lambda item: item.row):
            writer.writerow(result.__dict__)


def write_failed_manifest(path: Path, results: list[DownloadResult]) -> None:
    failed_results = [result for result in results if result.status == "failed"]
    write_manifest(path, failed_results)


def print_dry_run(tasks: list[DownloadTask]) -> None:
    print(f"Planned downloads: {len(tasks)}")
    for task in tasks[:10]:
        print(f"row={task.row} audio_id={task.audio_id} -> {task.output_path}")
    if len(tasks) > 10:
        print(f"... {len(tasks) - 10} more")


def run_downloads(args: argparse.Namespace, tasks: list[DownloadTask]) -> list[DownloadResult]:
    results: list[DownloadResult] = []
    counts = {"downloaded": 0, "exists": 0, "failed": 0}
    workers = max(1, args.workers)

    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = [
            executor.submit(
                download_one,
                task,
                timeout=args.timeout,
                retries=args.retries,
                overwrite=args.overwrite,
            )
            for task in tasks
        ]

        if tqdm is not None:
            progress = tqdm(total=len(futures), unit="file", dynamic_ncols=True)
        else:
            progress = None
            print(f"Progress: 0/{len(futures)}")

        try:
            for done_count, future in enumerate(as_completed(futures), start=1):
                result = future.result()
                results.append(result)
                counts[result.status] = counts.get(result.status, 0) + 1

                if result.status == "failed":
                    message = (
                        f"FAILED row={result.row} audio_id={result.audio_id}: "
                        f"{result.error}"
                    )
                    if progress is not None:
                        tqdm.write(message)
                    else:
                        print(message)

                if progress is not None:
                    progress.set_postfix(
                        downloaded=counts.get("downloaded", 0),
                        exists=counts.get("exists", 0),
                        failed=counts.get("failed", 0),
                    )
                    progress.update(1)
                elif done_count % 20 == 0 or done_count == len(futures):
                    print(
                        f"Progress: {done_count}/{len(futures)} "
                        f"downloaded={counts.get('downloaded', 0)} "
                        f"exists={counts.get('exists', 0)} "
                        f"failed={counts.get('failed', 0)}",
                        flush=True,
                    )
        finally:
            if progress is not None:
                progress.close()

    return results


def main() -> int:
    args = parse_args()
    args.excel = args.excel.expanduser().resolve()
    args.out_dir = args.out_dir.expanduser()
    manifest = args.manifest or (args.out_dir / "download_manifest.csv")
    failed_manifest = args.failed_manifest or (args.out_dir / "failed_downloads.csv")

    tasks = load_tasks(args)
    print(f"Excel: {args.excel}")
    print(f"Output directory: {args.out_dir}")
    print(f"Tasks: {len(tasks)}")

    if args.dry_run:
        print_dry_run(tasks)
        return 0
    if not tasks:
        print("No download tasks found.")
        return 0

    results = run_downloads(args, tasks)
    write_manifest(manifest, results)
    write_failed_manifest(failed_manifest, results)

    downloaded = sum(1 for item in results if item.status == "downloaded")
    exists = sum(1 for item in results if item.status == "exists")
    failed = sum(1 for item in results if item.status == "failed")
    print(
        "Done: "
        f"downloaded={downloaded} exists={exists} failed={failed} "
        f"manifest={manifest} failed_manifest={failed_manifest}"
    )
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
